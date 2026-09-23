"""Builds the OE testing working paper as an Excel file."""
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


GREEN = "05714D"
WHITE = "FFFFFF"
BLACK = "000000"
ARIAL = "Arial"


def write_working_paper(
        output_path: Path,
        metadata: dict[str, Any],
        overall_conclusion: str,
        sample_rows: list[dict[str, Any]],
) -> None:
    """Writes the metadata header block and the per-sample results table to one sheet."""
    metadata_rows = dict(metadata)
    metadata_rows["Overall Conclusion"] = overall_conclusion
    metadata_rows["Confirm that a human has reviewed the outcome of the assisted testing performed here"] = (
        "<to be filled by the reviewer manually>"
    )
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df = pd.DataFrame(sample_rows)
        metadata_start_row = 5
        sample_header_row = metadata_start_row + len(metadata_rows) + 2
        df.to_excel(writer, sheet_name="Working Paper", index=False, startrow=sample_header_row - 1)

        ws = writer.sheets["Working Paper"]
        max_table_columns = max(4, len(df.columns))
        title_end_column = get_column_letter(max_table_columns)
        thin_green = Side(style="thin", color=GREEN)
        table_border = Border(left=thin_green, right=thin_green, top=thin_green, bottom=thin_green)
        green_fill = PatternFill("solid", fgColor=GREEN)
        title_font = Font(name=ARIAL, bold=True, size=20, color=BLACK)
        section_font = Font(name=ARIAL, bold=True, size=14, color=BLACK)
        label_font = Font(name=ARIAL, bold=True, size=10, color=WHITE)
        body_font = Font(name=ARIAL, size=10, color=BLACK)
        centered = Alignment(horizontal="center", vertical="center")
        wrapped = Alignment(vertical="top", wrap_text=True)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_table_columns)
        ws.cell(row=1, column=1, value="Group Internal Audit").font = title_font
        ws.cell(row=1, column=1).alignment = centered
        ws.row_dimensions[1].height = 28

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_table_columns)
        ws.cell(row=3, column=1, value="Test of Effectiveness Workpaper").font = section_font
        ws.cell(row=3, column=1).alignment = centered
        ws.row_dimensions[3].height = 22

        for offset, (key, value) in enumerate(metadata_rows.items()):
            row = metadata_start_row + offset
            label_cell = ws.cell(row=row, column=1, value=key)
            label_cell.font = label_font
            label_cell.fill = green_fill
            label_cell.border = table_border
            label_cell.alignment = wrapped

            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            value_cell = ws.cell(row=row, column=2, value=value)
            value_cell.font = body_font
            value_cell.border = table_border
            value_cell.alignment = wrapped
            for column in range(2, 5):
                ws.cell(row=row, column=column).border = table_border
            ws.row_dimensions[row].height = 60 if key in {"Test Objective", "Overall Conclusion", "Human Review Confirmation"} else 20

        sample_end_row = sample_header_row + len(df.index)
        for row in ws.iter_rows(min_row=sample_header_row, max_row=sample_end_row, min_col=1, max_col=max_table_columns):
            for cell in row:
                cell.border = table_border
                cell.alignment = wrapped
                cell.font = label_font if cell.row == sample_header_row else body_font
                if cell.row == sample_header_row:
                    cell.fill = green_fill
        ws.row_dimensions[sample_header_row].height = 24

        review_row = sample_end_row + 2
        review_label = ws.cell(row=review_row, column=1, value="Human Review Confirmation")
        review_label.font = label_font
        review_label.fill = green_fill
        review_label.border = table_border
        review_label.alignment = wrapped
        ws.merge_cells(start_row=review_row, start_column=2, end_row=review_row, end_column=max_table_columns)
        review_value = ws.cell(
            row=review_row,
            column=2,
            value=metadata_rows.get("Human Review Confirmation", ""),
        )
        review_value.font = body_font
        review_value.alignment = wrapped
        review_value.border = table_border
        for column in range(2, max_table_columns + 1):
            ws.cell(row=review_row, column=column).border = table_border
        ws.row_dimensions[review_row].height = 30

        ws.column_dimensions["A"].width = 28
        for column in range(2, 5):
            ws.column_dimensions[get_column_letter(column)].width = 28
        for column in range(5, max_table_columns + 1):
            ws.column_dimensions[get_column_letter(column)].width = 18
        ws.freeze_panes = f"A{sample_header_row + 1}"
