"""Non-AI, format/regex-based redaction of Office documents and PDFs.

Rules are user-defined "format templates" (e.g. ``999-999-999``) or raw regular
expressions, matched against document text and replaced with either a masking
string (``xxxxxxxxx``) or a descriptive label (``<TFN>``). Each rule is sense-checked
before use so it can't accidentally redact unrelated, generic text.
"""
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

FORMAT_TOKEN_MAP = {"9": r"\d", "A": r"[A-Z]", "a": r"[a-z]", "X": r"[A-Za-z]", "*": r"[A-Za-z0-9]"}

# Benign sample sentences used to sense-check that a rule isn't too generic.
_GENERIC_SAMPLE_TEXT = [
    "the quick brown fox jumps over the lazy dog",
    "Hello World, this is a test document",
    "Invoice Number 12345 dated 2024-01-01",
    "Total: $1,234.56 due on receipt",
    "Meeting Notes - Project Kickoff",
    "Page 1 of 10",
    "Confidential Draft Version 2.0",
    "Annual Report 2023 Financial Summary",
    "Room 101, Section 4.2, Chapter Three",
    "Reference ABC-123-XYZ Item #4567",
    "Employee ID 000111 started on Monday",
    "Contact us at support for more information",
    "Phone: 0400 123 456",
    "Order 987654 shipped on schedule",
]

_TOO_GENERIC_PATTERNS = {
    ".*", ".+", "^.*$", "^.+$", r"\w+", r"\S+", r".", r"\d", r"[a-zA-Z]", r"\w", r"\d+", r"[a-z]+", r"[A-Z]+",
}


@dataclass
class RedactionRule:
    label: str
    pattern_type: str  # "format" or "regex"
    pattern: str
    mode: str  # "mask" or "label"
    regex: re.Pattern


_IPV4_OCTET = r"(?:25[0-5]|2[0-4]\d|1?\d?\d)"

# Fixed-format, vetted patterns. Order matters: more specific patterns run first so e.g. a URL
# containing an e-mail-like string is redacted as a whole.
PRESET_PATTERNS: list[dict[str, str]] = [
    {"id": "jwt", "label": "JWT_TOKEN", "name": "JWT tokens",
     "regex": r"\beyJ[A-Za-z0-9_-]{5,}\.[A-Za-z0-9_-]{5,}\.[A-Za-z0-9_-]*"},
    {"id": "secret", "label": "SECRET_KEY", "name": "Secrets / API keys",
     "regex": r"\b(?:AKIA|ASIA)[A-Z0-9]{16}\b|\bgh[pousr]_[A-Za-z0-9]{36,}\b|\bsk-[A-Za-z0-9_-]{20,}"
              r"|(?i:\b(?:api[_-]?key|secret(?:[_-]?key)?|access[_-]?token|auth[_-]?token|password|passwd|pwd)\b"
              r"\s*[:=]\s*[^\s,;]+)|(?i:\bbearer\s+[A-Za-z0-9._~+/=-]{16,})"},
    {"id": "url", "label": "URL", "name": "URLs",
     "regex": r"\b(?:https?|ftp)://[^\s<>\"')\]]+|\bwww\.[^\s<>\"')\]]+"},
    {"id": "email", "label": "EMAIL_ADDRESS", "name": "Email addresses",
     "regex": r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}"},
    {"id": "iban", "label": "IBAN", "name": "IBAN (international bank account number)",
     "regex": r"\b[A-Z]{2}\d{2}(?: ?[A-Z0-9]{4}){2,7}(?: ?[A-Z0-9]{1,3})?\b"},
    {"id": "au_bank", "label": "AU_BANK_ACCOUNT", "name": "Australian BSB + account number (e.g. 062-000 12345678)",
     "regex": r"(?<!\d)\d{3}[- ]?\d{3}[ -]\d{6,10}(?!\d)"},
    {"id": "au_bsb", "label": "AU_BSB", "name": "Australian BSB (e.g. 062-000)",
     "regex": r"(?<!\d)\d{3}-\d{3}(?![\d-])"},
    {"id": "card", "label": "CREDIT_CARD", "name": "Credit card numbers (Visa, Mastercard, Amex, Discover)",
     "regex": r"(?<!\d)(?:4\d{3}|5[1-5]\d{2}|2[2-7]\d{2}|6011|65\d{2})(?:[ -]?\d{4}){3}(?!\d)"
              r"|(?<!\d)3[47]\d{2}[ -]?\d{6}[ -]?\d{5}(?!\d)"},
    {"id": "abn", "label": "ABN", "name": "Australian Business Number (e.g. 12 345 678 901)",
     "regex": r"(?<!\d)\d{2} \d{3} \d{3} \d{3}(?!\d)"},
    {"id": "ipv4", "label": "IP_ADDRESS", "name": "IPv4 addresses",
     "regex": rf"(?<![\d.])(?:{_IPV4_OCTET}\.){{3}}{_IPV4_OCTET}(?![\d.])"},
    {"id": "ipv6", "label": "IPV6_ADDRESS", "name": "IPv6 addresses",
     "regex": r"(?<![:\w])(?:(?:[A-Fa-f0-9]{1,4}:){7}[A-Fa-f0-9]{1,4}"
              r"|(?:[A-Fa-f0-9]{1,4}:){1,7}:(?:[A-Fa-f0-9]{1,4}(?::[A-Fa-f0-9]{1,4}){0,6})?"
              r"|::(?:[A-Fa-f0-9]{1,4}(?::[A-Fa-f0-9]{1,4}){0,6}))(?![:\w])"},
    {"id": "mac", "label": "MAC_ADDRESS", "name": "MAC addresses",
     "regex": r"(?<![0-9A-Fa-f:-])(?:[0-9A-Fa-f]{2}[:-]){5}[0-9A-Fa-f]{2}(?![0-9A-Fa-f:-])"},
    {"id": "uuid", "label": "UUID", "name": "UUIDs / GUIDs",
     "regex": r"\b[0-9A-Fa-f]{8}-[0-9A-Fa-f]{4}-[0-9A-Fa-f]{4}-[0-9A-Fa-f]{4}-[0-9A-Fa-f]{12}\b"},
]
_PRESETS_BY_ID = {p["id"]: p for p in PRESET_PATTERNS}


def format_to_regex(fmt: str) -> str:
    """Translates a 9/A/a/X/* placeholder format into a word-bounded regex."""
    parts = [FORMAT_TOKEN_MAP.get(ch, re.escape(ch)) for ch in fmt]
    return r"(?<![A-Za-z0-9])" + "".join(parts) + r"(?![A-Za-z0-9])"


def validate_rule(label: str, pattern_type: str, pattern: str) -> list[str]:
    """Sense-checks one redaction rule, returning a list of problems (empty = OK)."""
    problems: list[str] = []
    if not label.strip():
        problems.append("Every redaction rule needs a label (e.g. TFN, PHONE_NUMBER).")

    stripped = pattern.strip()
    if not stripped:
        problems.append("Pattern cannot be empty.")
        return problems

    if pattern_type == "format":
        if len(stripped) < 4:
            problems.append(f"Format '{pattern}' is too short to safely target specific data (minimum 4 characters).")
        literal_count = sum(1 for c in stripped if c not in FORMAT_TOKEN_MAP)
        if literal_count == 0 and len(stripped) < 5:
            problems.append(
                f"Format '{pattern}' has no fixed/literal characters and is too short/generic "
                "(use at least 5 wildcard positions, or add literal separators)."
            )
        regex_str = format_to_regex(stripped)
    else:
        if stripped in _TOO_GENERIC_PATTERNS:
            problems.append(f"Pattern '{pattern}' is too generic and would match almost any text.")
        regex_str = stripped

    try:
        compiled = re.compile(regex_str)
    except re.error as exc:
        problems.append(f"Pattern '{pattern}' is not a valid regular expression: {exc}")
        return problems

    if compiled.match(""):
        problems.append(f"Pattern '{pattern}' matches an empty string, which is too permissive.")
        return problems

    hits = sum(1 for sample in _GENERIC_SAMPLE_TEXT if compiled.search(sample))
    if hits > len(_GENERIC_SAMPLE_TEXT) * 0.35:
        problems.append(
            f"Pattern '{pattern}' matched {hits}/{len(_GENERIC_SAMPLE_TEXT)} unrelated sample sentences during "
            "a sense check - it looks too generic and could redact unrelated information. Please make it more specific."
        )
    return problems


def build_rules(raw_rules: list[dict]) -> tuple[list[RedactionRule], list[str]]:
    """Validates and compiles a list of {label, pattern_type, pattern, mode} dicts."""
    rules: list[RedactionRule] = []
    errors: list[str] = []
    for r in raw_rules:
        preset = _PRESETS_BY_ID.get(str(r.get("preset", "")))
        if preset:
            mode = r.get("mode", "mask")
            rules.append(RedactionRule(
                label=preset["label"], pattern_type="regex", pattern=preset["regex"],
                mode=mode if mode in {"mask", "label"} else "mask", regex=re.compile(preset["regex"]),
            ))
            continue
        if r.get("preset"):
            errors.append(f"Unknown preset '{r['preset']}'.")
            continue
        label = str(r.get("label", "")).strip()
        pattern_type = r.get("pattern_type", "format")
        pattern = str(r.get("pattern", "")).strip()
        mode = r.get("mode", "mask")
        problems = validate_rule(label, pattern_type, pattern)
        if problems:
            errors.extend(problems)
            continue
        regex_str = format_to_regex(pattern) if pattern_type == "format" else pattern
        rules.append(RedactionRule(label=label, pattern_type=pattern_type, pattern=pattern, mode=mode, regex=re.compile(regex_str)))
    return rules, errors


def _replacement_for(rule: RedactionRule, matched_text: str) -> str:
    return "x" * len(matched_text) if rule.mode == "mask" else f"<{rule.label}>"


def _redact_text(text: str, rules: list[RedactionRule], counts: dict[str, int]) -> str:
    for rule in rules:
        def _sub(m: re.Match, rule=rule) -> str:
            counts[rule.label] = counts.get(rule.label, 0) + 1
            return _replacement_for(rule, m.group(0))
        text = rule.regex.sub(_sub, text)
    return text


def redact_docx(src: Path, dest: Path, rules: list[RedactionRule]) -> dict[str, int]:
    from docx import Document as DocxDocument
    doc = DocxDocument(str(src))
    counts: dict[str, int] = {}

    def _process_paragraphs(paragraphs):
        for p in paragraphs:
            for run in p.runs:
                if run.text:
                    run.text = _redact_text(run.text, rules, counts)

    _process_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                _process_paragraphs(cell.paragraphs)
    for section in doc.sections:
        for header_footer in (section.header, section.footer):
            _process_paragraphs(header_footer.paragraphs)
    doc.save(str(dest))
    return counts


def redact_pptx(src: Path, dest: Path, rules: list[RedactionRule]) -> dict[str, int]:
    from pptx import Presentation
    prs = Presentation(str(src))
    counts: dict[str, int] = {}

    def _process_text_frame(text_frame):
        for p in text_frame.paragraphs:
            for run in p.runs:
                if run.text:
                    run.text = _redact_text(run.text, rules, counts)

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                _process_text_frame(shape.text_frame)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        _process_text_frame(cell.text_frame)
    prs.save(str(dest))
    return counts


def redact_xlsx(src: Path, dest: Path, rules: list[RedactionRule]) -> dict[str, int]:
    import openpyxl
    wb = openpyxl.load_workbook(str(src))
    counts: dict[str, int] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value:
                    cell.value = _redact_text(cell.value, rules, counts)
    wb.save(str(dest))
    return counts


def redact_pdf(src: Path, dest: Path, rules: list[RedactionRule]) -> dict[str, int]:
    """Redacts matching whole "words" in a PDF by covering them with a box + replacement text.

    Uses no AI/OCR, so only matches contained within a single extracted word token are
    found - patterns spanning whitespace (e.g. multi-word phrases) will not be redacted.
    """
    import pymupdf as fitz
    counts: dict[str, int] = {}
    doc = fitz.open(str(src))
    for page in doc:
        for w in page.get_text("words"):
            token = w[4]
            for rule in rules:
                if rule.regex.search(token):
                    counts[rule.label] = counts.get(rule.label, 0) + 1
                    rect = fitz.Rect(w[0], w[1], w[2], w[3])
                    replacement = _replacement_for(rule, token)
                    fill = (0, 0, 0) if rule.mode == "mask" else (1, 1, 0.6)
                    text_color = (1, 1, 1) if rule.mode == "mask" else (0, 0, 0)
                    page.add_redact_annot(rect, text=replacement, fill=fill, text_color=text_color, fontsize=8)
                    break
        page.apply_redactions()
    doc.save(str(dest))
    return counts


_HANDLERS: dict[str, Callable[[Path, Path, list[RedactionRule]], dict[str, int]]] = {
    ".docx": redact_docx,
    ".pptx": redact_pptx,
    ".xlsx": redact_xlsx,
    ".pdf": redact_pdf,
}


def redact_files(file_paths: list[Path], rules: list[RedactionRule], output_dir: Path) -> tuple[Path, list[dict]]:
    """Redacts each file, zips the redacted outputs into output_dir, and returns (zip_path, summary)."""
    output_dir.mkdir(parents=True, exist_ok=True)
    summary: list[dict] = []
    redacted_paths: list[Path] = []
    for src in file_paths:
        ext = src.suffix.lower()
        handler = _HANDLERS.get(ext)
        if handler is None:
            summary.append({"filename": src.name, "ok": False, "error": f"Unsupported file type '{ext}'.", "redacted_counts": {}})
            continue
        dest = output_dir / f"redacted_{src.name}"
        try:
            counts = handler(src, dest, rules)
            redacted_paths.append(dest)
            summary.append({"filename": src.name, "ok": True, "error": None, "redacted_counts": counts})
        except Exception as exc:  # noqa: BLE001 - surface per-file failure, keep processing the rest
            summary.append({"filename": src.name, "ok": False, "error": str(exc), "redacted_counts": {}})

    zip_path = output_dir / "redacted_documents.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in redacted_paths:
            zf.write(p, arcname=p.name)
    return zip_path, summary
